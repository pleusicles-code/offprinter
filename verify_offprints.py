#!/usr/bin/env python3
"""Verify OFFPRINTER metadata using Crossref, OpenAlex and selective OpenAI comparison."""
import argparse, json, os, re, sys, time, unicodedata
from pathlib import Path
from difflib import SequenceMatcher
from urllib.request import Request, urlopen
from urllib.parse import urlencode
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
CACHE = ROOT / ".offprinter_verify_cache.json"

def norm(x):
    x = unicodedata.normalize("NFKD", str(x or "").casefold())
    return " ".join(re.sub(r"[^\w]+", " ", "".join(c for c in x if not unicodedata.combining(c))).split())

def sim(a,b):
    a,b=norm(a),norm(b)
    if not a or not b: return 0.
    return max(SequenceMatcher(None,a,b).ratio(),
               SequenceMatcher(None," ".join(sorted(a.split()))," ".join(sorted(b.split()))).ratio())

def family_name(normalized_author):
    """OFFPRINTER author_normalized begins with the family name (with or without comma)."""
    first = str(normalized_author or "").split(";")[0].strip()
    return first.split(",", 1)[0].split()[0] if first else ""


def author_sim(a, b):
    """Compare a catalogue family name against all words in a repository author name."""
    family = family_name(a)
    candidates = str(b or "").split(";")
    if not family or not b:
        return 0.0
    family_score = max(
        (sim(family, word) for person in candidates
         for word in re.findall(r"[^\\W_]+(?:[-'][^\\W_]+)*", person, re.UNICODE)),
        default=0.0,
    )
    # The surname is the stable identifier across Hungarian / Western name order
    # and translated or abbreviated given names. Do not allow a fuzzy full-name
    # score to override an absent surname.
    full_score = max(
        (sim(str(a).split(";")[0].replace(",", " "), person.replace(",", " "))
         for person in candidates),
        default=0.0,
    )
    return 0.85 * family_score + 0.15 * full_score

def value(v): return str(v).strip() if v is not None else ""

def request(url,params,cache):
    endpoint=url+"?"+urlencode(params)
    if endpoint in cache: return cache[endpoint]
    for attempt in range(3):
        try:
            req=Request(endpoint,headers={"User-Agent":"OFFPRINTER/1.0 bibliographic research","Accept":"application/json"})
            with urlopen(req,timeout=20) as response:
                obj=json.loads(response.read().decode("utf-8"))
            cache[endpoint]=obj
            return obj
        except Exception:
            if attempt==2: raise
            time.sleep(2*(attempt+1))

def crossref(title,cache):
    data=request("https://api.crossref.org/works",{"query.bibliographic":title,"rows":8},cache)
    for item in data.get("message",{}).get("items",[]):
        dp=(item.get("published-print") or item.get("published-online") or item.get("published") or {}).get("date-parts",[])
        yield dict(source="Crossref",title=(item.get("title") or [""])[0],
            author="; ".join(" ".join(x for x in (a.get("given",""),a.get("family","")) if x) for a in item.get("author",[])),
            publication=(item.get("container-title") or [""])[0],year=str(dp[0][0]) if dp and dp[0] else "",
            volume=value(item.get("volume")),issue=value(item.get("issue")),pages=value(item.get("page")),
            url=item.get("URL") or ("https://doi.org/"+item.get("DOI","")))

def openalex(title,cache,key):
    params={"search":title,"per_page":10}
    if key: params["api_key"]=key
    data=request("https://api.openalex.org/works",params,cache)
    for item in data.get("results",[]):
        loc=item.get("primary_location") or {}; journal=(loc.get("source") or {}).get("display_name","")
        bib=item.get("biblio") or {}
        first,last=bib.get("first_page"),bib.get("last_page")
        links=[x.get("landing_page_url","") for x in item.get("locations",[]) or []]
        real=next((x for x in links if "real.mtak.hu" in x or "real-j.mtak.hu" in x),"")
        yield dict(source="OpenAlex"+(" / REAL" if real else ""),title=item.get("title") or "",
            author="; ".join((x.get("author") or {}).get("display_name","") for x in item.get("authorships",[])),
            publication=journal,year=value(item.get("publication_year")),
            volume=value(bib.get("volume")),issue=value(bib.get("issue")),
            pages=str(first)+"-"+str(last) if first and last else "",
            url=real or loc.get("landing_page_url") or item.get("id",""))

def rank(rec,item):
    t=sim(rec.get("Title"),item["title"])
    a=author_sim(rec.get("Author_normalized") or rec.get("Author"),item["author"])
    j=sim(rec.get("Publication"),item["publication"])
    return (.76*t+.19*a+.05*j,t,a,j)

def ai_verdict(rec,candidate,key):
    from openai import OpenAI
    client=OpenAI(api_key=key,max_retries=2,timeout=45)
    prompt=("Assess whether these are the SAME publication, tolerating OCR mistakes, name order, "
            "journal abbreviations, and errors in year/volume. Return ONLY SAME, DIFFERENT or UNCERTAIN. "
            "Do not invent evidence.\n"+json.dumps({"input":rec,"candidate":candidate},ensure_ascii=False,default=str))
    result=client.responses.create(model="gpt-5.4-mini",reasoning={"effort":"none"},store=False,
                                   input=prompt)
    text=(result.output_text or "").strip().upper()
    return text if text in ("SAME","DIFFERENT","UNCERTAIN") else "UNCERTAIN"

def verify(rec,cache,key,openalex_key,ai_mode):
    candidates=[];errors=[]
    for provider in (lambda:list(crossref(value(rec.get("Title")),cache)),
                     lambda:list(openalex(value(rec.get("Title")),cache,openalex_key))):
        try: candidates+=provider()
        except Exception as e: errors.append(str(e)[:85])
    if not candidates:
        return {},("SEARCH ERROR: "+"; ".join(errors) if errors else "NOT FOUND"),""
    ranked=sorted([(rank(rec,c),c) for c in candidates],key=lambda x:x[0][0],reverse=True)
    (overall,t,a,j),best=ranked[0]
    if t<.90 or (value(rec.get("Author")) and a<.76) or overall<.88:
        return {},"UNVERIFIED (no strong author/title match)",""
    if len(ranked)>1 and overall-ranked[1][0][0]<.035:
        other=ranked[1][1]
        same=sim(best["title"],other["title"])>.97 and author_sim(best["author"],other["author"])>.9
        if not same: return {},"AMBIGUOUS (several plausible records)",best["url"]
    if key and ai_mode!="never" and (ai_mode=="always" or t<.97 or a<.94 or (j<.75 and value(rec.get("Publication")))):
        try: verdict=ai_verdict(rec,best,key)
        except Exception as e: return {},"UNVERIFIED (AI comparison failed)",best["url"]
        if verdict!="SAME": return {},"UNVERIFIED (AI: "+verdict+")",best["url"]
    if t<.96 or overall<.94 or (value(rec.get("Author")) and a<.85):
        return {},"POSSIBLE MATCH (not safe to change)",best["url"]
    changes={}; notes=[]; warnings=[]
    title=value(rec.get("Title"))
    cleaned=re.sub(r"[\s.;:,]+$","",title).strip()
    if title!=cleaned and cleaned:
        changes["Title"]=cleaned; notes.append("Title: removed terminal punctuation")
    for keyname,ext in (("Year","year"),("Volume","volume"),("Issue","issue"),("Pages","pages")):
        old=value(rec.get(keyname));new=value(best.get(ext))
        if not new or old==new or norm(old)==norm(new): continue
        # Two agreeing strong records can corroborate a correction; conflicting evidence blocks it.
        support=[x for _,x in ranked if sim(rec.get("Title"),x["title"])>=.94
                 and author_sim(rec.get("Author"),x["author"])>=.80 and value(x.get(ext))]
        distinct={norm(x[ext]) for x in support}
        if len(distinct)>1:
            warnings.append(keyname+" conflicts across sources");continue
        if keyname=="Year" and not re.fullmatch(r"(1[5-9]\d\d|20\d\d)",new):
            warnings.append("Year not a single date");continue
        if keyname=="Pages" and not re.fullmatch(r"\d+[A-Za-z]?[-–]\d+[A-Za-z]?",new):
            continue
        changes[keyname]=new
        notes.append(keyname+": "+(old or "∅")+" → "+new)
    if warnings:
        return changes,"CHECK: "+"; ".join(notes+warnings),best["url"]
    if not notes:
        compared=[k for k,ext in (("Year","year"),("Volume","volume"),("Issue","issue"),("Pages","pages"))
                  if value(rec.get(k)) and value(best.get(ext))]
        if not compared:
            return {},"MATCHED (bibliographic numbering not independently verified)",best["url"]
        missing=[k for k,ext in (("Year","year"),("Volume","volume"),("Issue","issue"),("Pages","pages"))
                 if value(rec.get(k)) and not value(best.get(ext))]
        if missing:
            return {},"PARTIALLY VERIFIED (source lacks "+", ".join(missing)+")",best["url"]
        return {},"VERIFIED CORRECT",best["url"]
    if all("removed terminal punctuation" in n for n in notes):
        return changes,"VERIFIED (CLEANUP): "+"; ".join(notes),best["url"]
    return changes,"VERIFIED (CORRECTED): "+"; ".join(notes),best["url"]

def main():
    ap=argparse.ArgumentParser(description=__doc__)
    ap.add_argument("xlsx",type=Path)
    ap.add_argument("--limit",type=int,default=None)
    ap.add_argument("--ai",choices=("ambiguous","always","never"),default="ambiguous")
    ap.add_argument("--openalex-key",default=os.getenv("OPENALEX_API_KEY",""))
    ap.add_argument("--no-cache",action="store_true")
    args=ap.parse_args()
    path=args.xlsx.resolve()
    if not path.exists(): ap.error("Workbook not found")
    wb=load_workbook(path)
    if "Catalogue" not in wb: ap.error("Catalogue worksheet missing")
    sh=wb["Catalogue"]
    cols={value(c.value):i for i,c in enumerate(sh[1],1) if c.value is not None}
    for required in ("Title","Author","Publication"):
        if required not in cols: ap.error("Missing column: "+required)
    logcol=cols.get("Verification") or sh.max_column+1
    sh.cell(1,logcol,"Verification")
    sourcecol=cols.get("Verification source") or max(logcol,sh.max_column)+1
    sh.cell(1,sourcecol,"Verification source")
    cache={}
    if not args.no_cache and CACHE.exists():
        try: cache=json.loads(CACHE.read_text(encoding="utf-8"))
        except Exception: pass
    keypath=ROOT/".openai_api_key"
    key=keypath.read_text(encoding="utf-8").strip() if keypath.exists() else None
    if args.ai!="never" and not key: print("OpenAI key not found; AI comparisons disabled",file=sys.stderr)
    count=0
    for r in range(2,sh.max_row+1):
        if args.limit and count>=args.limit: break
        rec={k:sh.cell(r,c).value for k,c in cols.items()}
        try:
            changes,status,link=verify(rec,cache,key,args.openalex_key,args.ai) if value(rec.get("Title")) else ({}, "UNVERIFIED (no title)","")
        except Exception as e:
            changes,status,link={},"SEARCH ERROR: "+str(e)[:120],""
        for k,v in changes.items():
            if k in cols: sh.cell(r,cols[k],v)
        sh.cell(r,logcol,status)
        sc=sh.cell(r,sourcecol,link)
        if link: sc.hyperlink=link;sc.style="Hyperlink"
        count+=1
        print(f"[{count}] row {r}: {status}",flush=True)
        if not args.no_cache and count%10==0:
            CACHE.write_text(json.dumps(cache,ensure_ascii=False),encoding="utf-8")
    if not args.no_cache: CACHE.write_text(json.dumps(cache,ensure_ascii=False),encoding="utf-8")
    for col in (logcol,sourcecol):
        cell=sh.cell(1,col)
        cell.fill=PatternFill("solid",fgColor="23445A")
        cell.font=Font(bold=True,color="FFFFFF")
    sh.column_dimensions[get_column_letter(logcol)].width=74
    sh.column_dimensions[get_column_letter(sourcecol)].width=44
    sh.auto_filter.ref=sh.dimensions
    dest=path.with_name(path.stem+"_verified.xlsx")
    wb.save(dest)
    print(f"Wrote {dest}: {count} rows checked")

if __name__=="__main__": main()
